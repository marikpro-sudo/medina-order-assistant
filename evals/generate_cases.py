"""Deterministic eval case generator (SPEC section 1, seed 7)."""

from __future__ import annotations

import argparse
import csv
import json
import random
import re
from pathlib import Path

import openpyxl
import xlwt

from data_access.search import product_form_family
from data_access.synthetic import default_source
from evals.render_photo import render_list

SEED = 7

N_TEXT_RU = 40
N_TEXT_KK = 15
N_SHEET = 20
N_PORTAL = 10
N_PHOTO = 10
N_LONG = 5

ERROR_CODES = (
    "not_text",
    "bad_ean",
    "bad_material_id",
    "ean_code_mismatch",
    "bad_qty",
    "not_multiple",
    "column_count",
    "duplicate",
)

UNRECOGNIZED = [
    "бахилы",
    "маски медицинские 50 шт",
    "шприц 5 мл",
    "перчатки нитриловые M",
    "лейкопластырь",
    "фывапр 10",
    "ывапрыв 5 шт",
    "кцукен 20 уп",
]

_RU_GREETINGS = ["Здравствуйте!", "Добрый день!", "Добрый день, коллеги!", ""]
_RU_CLOSERS = ["Спасибо!", "Ждем счет.", "Срочно, пожалуйста.", ""]
_KK_GREETINGS = ["Сәлеметсіз бе!", "Қайырлы күн!"]

_RU_QTY = ["{q}", "{q} уп", "x{q}", "{q}шт", "- {q}"]
_RU_QTY_PLAIN = ["{q}", "{q} уп", "{q} шт"]
_KK_QTY = ["{q} қорап", "{q} дана", "{q} шт", "{q} уп"]

_RU_LAT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "yo",
    "ж": "zh", "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m",
    "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u",
    "ф": "f", "х": "kh", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "shch",
    "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
}

_FORM_SHORTHAND = {
    "tab": "таб", "caps": "капс", "syrup": "сироп", "susp": "сусп",
    "drops": "капли", "spray": "спрей", "ointment": "мазь", "cream": "крем",
    "gel": "гель", "supp": "супп", "inj": "амп", "powder": "пор",
    "solution": "р-р",
}

_SHEET_LAYOUTS = [
    {"headers": ["Наименование", "Код товара", "Количество"], "title": None, "numbered": False},
    {"headers": ["Товар", "Артикул", "Кол-во"], "title": "Заказ от аптеки", "numbered": True},
    {"headers": ["Атауы", "Тауар коды", "Саны"], "title": "Тапсырыс", "numbered": False},
    {"headers": ["Product", "Material", "Qty"], "title": "Order", "numbered": True},
    {"headers": ["Штрихкод", "Код", "шт"], "title": None, "numbered": False},
]
_SHEET_EXTS = [".xlsx", ".xls", ".csv"]


def transliterate(text: str) -> str:
    out = []
    for ch in text:
        low = ch.lower()
        if low in _RU_LAT:
            lat = _RU_LAT[low]
            out.append(lat.capitalize() if ch.isupper() else lat)
        else:
            out.append(ch)
    return "".join(out)


def typo_word(rng: random.Random, word: str) -> str:
    if len(word) < 6:
        return word
    op = rng.choice(["drop", "swap", "dup"])
    i = rng.randint(1, len(word) - 1)
    if op == "drop":
        return word[:i] + word[i + 1:]
    if op == "swap":
        return word[:i - 1] + word[i] + word[i - 1] + word[i + 1:]
    return word[:i] + word[i] + word[i:]


def noisify_name(rng: random.Random, name: str) -> str:
    if rng.random() < 0.10:
        name = transliterate(name)
    parts = []
    for tok in name.split():
        if tok.isalpha() and len(tok) >= 5 and rng.random() < 0.35:
            tok = typo_word(rng, tok)
        parts.append(tok)
    return " ".join(parts)


def qty_token(rng: random.Random, qty: int, kk: bool, plain: bool = False) -> str:
    if kk:
        style = rng.choice(_KK_QTY)
    elif plain:
        style = rng.choice(_RU_QTY_PLAIN)
    else:
        style = rng.choice(_RU_QTY)
    return style.format(q=qty)


def build_context(ds):
    """Precompute catalog-order candidate pools and per-branch stock."""
    products = list(ds.products())
    stock: dict[str, dict[str, int]] = {}
    for p in products:
        per_plant = {}
        for plant in ("P01", "P02", "P03"):
            per_plant[plant] = sum(b.qty_available for b in ds.batches(p.material_id, plant))
        stock[p.material_id] = per_plant
    families = {p.material_id: product_form_family(p.dosage_form) for p in products}
    by_inn: dict[str, list] = {}
    for p in products:
        by_inn.setdefault(p.inn, []).append(p)
    clarify_inns = sorted(
        inn for inn, plist in by_inn.items()
        if len({p.strength for p in plist}) >= 2
    )
    brand_keys: dict[str, set] = {}
    for p in products:
        brand_keys.setdefault(p.brand, set()).add((p.strength, p.dosage_form))
    brand_single = {brand: len(keys) == 1 for brand, keys in brand_keys.items()}

    def same_group(p, q) -> bool:
        return q.inn == p.inn and families[q.material_id] == families[p.material_id] \
            and q.strength == p.strength

    def pack_variants(p):
        return [q for q in products
                if q.material_id != p.material_id and q.brand == p.brand
                and q.strength == p.strength
                and families[q.material_id] == families[p.material_id]]

    def inn_group(p):
        return [q for q in products if same_group(p, q)]

    return {
        "products": products, "stock": stock, "families": families,
        "by_inn": by_inn, "clarify_inns": clarify_inns,
        "brand_single": brand_single, "same_group": same_group,
        "pack_variants": pack_variants, "inn_group": inn_group,
    }


def acceptable_found(ctx, gold, use_inn: bool, has_pack: bool) -> list[str]:
    acc = [gold.material_id]
    if not has_pack:
        acc += [q.material_id for q in ctx["pack_variants"](gold)]
    if use_inn:
        acc += [q.material_id for q in ctx["inn_group"](gold)]
    seen = set()
    out = []
    for mid in acc:
        if mid not in seen:
            seen.add(mid)
            out.append(mid)
    return out


def render_name_part(rng, ctx, p, use_inn: bool, brand_only: bool,
                     with_strength: bool, with_form: bool, has_pack: bool) -> str:
    name = noisify_name(rng, p.inn if use_inn else p.brand)
    parts = [name]
    if with_strength and not brand_only:
        parts.append(p.strength)
    if with_form and not brand_only:
        short = _FORM_SHORTHAND.get(ctx["families"][p.material_id], "")
        if short:
            parts.append(short)
    if has_pack:
        parts.append(p.pack)
    return " ".join(x for x in parts if x)


def pick_scenario(rng: random.Random) -> str:
    r = rng.random()
    if r < 0.60:
        return "found"
    if r < 0.75:
        return "rounded"
    if r < 0.90:
        return "analog"
    if r < 0.95:
        return "unrecognized"
    return "clarify"


def build_line(rng: random.Random, ctx, plant: str, kk: bool,
               plain_qty: bool = False):
    """Return (gold record, name part for spreadsheets, qty shown in text)."""
    stock = ctx["stock"]
    products = ctx["products"]
    scen = pick_scenario(rng)
    if scen == "rounded" and not any(
            p.order_multiple > 1 and stock[p.material_id][plant] >= p.order_multiple
            for p in products):
        scen = "found"
    if scen == "analog" and not any(
            stock[p.material_id][plant] == 0 and any(
                stock[q.material_id][plant] > 0 for q in ctx["inn_group"](p))
            for p in products):
        scen = "found"
    if scen == "clarify" and not ctx["clarify_inns"]:
        scen = "found"

    if scen == "unrecognized":
        text = rng.choice(UNRECOGNIZED)
        if not re.search(r"\d", text):
            text = f"{text} {rng.choice([5, 10, 20])} шт"
        gold = {"source": text, "expect": "unrecognized", "material_id": None,
                "acceptable": [], "gold_qty": 0}
        return gold, text, 0

    if scen == "clarify":
        inn = rng.choice(ctx["clarify_inns"])
        plist = ctx["by_inn"][inn]
        rep = plist[0]
        mult = rep.order_multiple
        qty = mult * rng.randint(1, 5)
        name = noisify_name(rng, inn)
        parts = [name]
        if rng.random() < 0.5:
            short = _FORM_SHORTHAND.get(ctx["families"][rep.material_id], "")
            if short:
                parts.append(short)
        token = qty_token(rng, qty, kk, plain_qty)
        source = " ".join(parts + [token])
        gold = {"source": source, "expect": "clarify", "material_id": rep.material_id,
                "acceptable": [p.material_id for p in plist], "gold_qty": qty}
        return gold, " ".join(parts), qty

    if scen == "analog":
        cands = [p for p in products
                 if stock[p.material_id][plant] == 0
                 and any(stock[q.material_id][plant] > 0 for q in ctx["inn_group"](p))]
        gold_p = rng.choice(cands)
        acc = [q for q in ctx["inn_group"](gold_p) if stock[q.material_id][plant] > 0]
        first = acc[0]
        qty = first.order_multiple * rng.randint(
            1, max(1, min(5, stock[first.material_id][plant] // first.order_multiple)))
        use_inn = rng.random() < 0.3
        has_pack = rng.random() < 0.25
        name_part = render_name_part(rng, ctx, gold_p, use_inn, False, True, True, has_pack)
        token = qty_token(rng, qty, kk, plain_qty)
        source = f"{name_part} {token}"
        gold = {"source": source, "expect": "analog", "material_id": gold_p.material_id,
                "acceptable": [q.material_id for q in acc], "gold_qty": qty}
        return gold, name_part, qty

    if scen == "rounded":
        cands = [p for p in products if p.order_multiple > 1
                 and stock[p.material_id][plant] >= p.order_multiple]
        gold_p = rng.choice(cands)
        mult = gold_p.order_multiple
        maxk = max(1, min(5, stock[gold_p.material_id][plant] // mult))
        k = rng.randint(1, maxk)
        delta = rng.randint(1, mult - 1)
        shown = k * mult - delta
        gold_qty = k * mult
    else:
        cands = [p for p in products if stock[p.material_id][plant] >= p.order_multiple]
        gold_p = rng.choice(cands)
        mult = gold_p.order_multiple
        maxk = max(1, min(5, stock[gold_p.material_id][plant] // mult))
        shown = mult * rng.randint(1, maxk)
        gold_qty = shown
    use_inn = rng.random() < 0.3
    has_pack = rng.random() < 0.25
    brand_only = (not use_inn and ctx["brand_single"].get(gold_p.brand, False)
                  and rng.random() < 0.3)
    if brand_only:
        has_pack = False
    name_part = render_name_part(rng, ctx, gold_p, use_inn, brand_only,
                                 True, True, has_pack)
    token = qty_token(rng, shown, kk, plain_qty)
    source = f"{name_part} {token}"
    gold = {"source": source, "expect": scen, "material_id": gold_p.material_id,
            "acceptable": acceptable_found(ctx, gold_p, use_inn, has_pack),
            "gold_qty": gold_qty}
    return gold, name_part, shown


def assemble_text_message(rng: random.Random, sources: list[str], kk: bool) -> str:
    if kk:
        greet = rng.choice(_KK_GREETINGS)
        sep = rng.choice(["\n", "\n", ", "])
        if len(sources) > 1 and sep == ", ":
            body = ", ".join(sources[:-1]) + " және " + sources[-1]
        elif len(sources) > 1:
            body = "\n".join(sources[:-1] + ["және " + sources[-1]])
        else:
            body = sources[0]
        parts = [greet, "Осы дәрілер керек:", body, "Рақмет!"]
    else:
        greet = rng.choice(_RU_GREETINGS)
        closer = rng.choice(_RU_CLOSERS)
        sep = rng.choice(["\n", "\n", "\n", ", ", "; "])
        body = sep.join(sources)
        parts = [greet, body, closer]
    return "\n".join(p for p in parts if p)


def write_sheet_file(path: Path, layout: dict, rows: list[tuple[str, str, str]]) -> None:
    headers = layout["headers"]
    numbered = layout["numbered"]
    title = layout["title"]
    body: list[list[str]] = []
    if title:
        body.append([title])
        body.append([])
    header_row = (["№"] if numbered else []) + headers
    body.append(header_row)
    for i, (name, code, qty) in enumerate(rows, start=1):
        body.append((([str(i)] if numbered else []) + [name, code, qty]))
    ext = path.suffix.lower()
    if ext == ".xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        for row in body:
            ws.append(row)
        wb.save(str(path))
    elif ext == ".xls":
        wb = xlwt.Workbook(encoding="utf-8")
        ws = wb.add_sheet("Order")
        text = xlwt.easyxf(num_format_str="@")
        for r, row in enumerate(body):
            for c, value in enumerate(row):
                ws.write(r, c, value, text)
        wb.save(str(path))
    elif ext == ".csv":
        with open(path, "w", encoding="cp1251", newline="") as f:
            writer = csv.writer(f, delimiter=";")
            writer.writerows(body)
    else:
        raise ValueError(f"unsupported sheet extension: {ext}")


def ean_bad(ean: str) -> str:
    bad_digit = (int(ean[-1]) + 1) % 10
    return ean[:-1] + str(bad_digit)


def write_portal_file(path: Path, rows: list[tuple[str, str, str, str]]) -> None:
    """Rows are (ean, material_id, qty, kind) with kind text|number|extra."""
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Order")
    text = xlwt.easyxf(num_format_str="@")
    for c, h in enumerate(("Штрихкод", "Код товара", "Количество")):
        ws.write(0, c, h, text)
    for r, (ean, mid, qty, kind) in enumerate(rows, start=1):
        if kind == "number":
            ws.write(r, 0, ean, text)
            ws.write(r, 1, mid, text)
            ws.write(r, 2, int(qty))
        elif kind == "extra":
            ws.write(r, 0, ean, text)
            ws.write(r, 1, mid, text)
            ws.write(r, 2, qty, text)
            ws.write(r, 3, "доп. колонка", text)
        else:
            ws.write(r, 0, ean, text)
            ws.write(r, 1, mid, text)
            ws.write(r, 2, qty, text)
    wb.save(str(path))


def make_sheet_case(rng, ctx, ds, plant, case_id, customer_id, seq, files_dir):
    n = rng.randint(4, 10)
    gold_lines = []
    sheet_rows = []
    for _ in range(n):
        gold, name_part, shown = build_line(rng, ctx, plant, kk=False, plain_qty=True)
        gold_lines.append(gold)
        if rng.random() < 0.25:
            p = ds.get_product(gold["material_id"]) if gold["material_id"] else None
            cell_name = p.ean13 if p is not None else name_part
            cell_code = ""
        else:
            cell_name = name_part
            if gold["material_id"] is not None and rng.random() < 0.5:
                cell_code = gold["material_id"]
            else:
                cell_code = ""
        token = qty_token(rng, shown, kk=False, plain=True) if shown else ""
        if gold["expect"] == "unrecognized":
            token = str(rng.randint(1, 20))
        sheet_rows.append((cell_name, cell_code, token))
        gold["source"] = f"{cell_name} {token}".strip()
    ext = _SHEET_EXTS[seq % len(_SHEET_EXTS)]
    layout = _SHEET_LAYOUTS[seq % len(_SHEET_LAYOUTS)]
    if ext == ".csv" and layout["headers"] != _SHEET_LAYOUTS[0]["headers"]:
        layout = _SHEET_LAYOUTS[0]
    if ext == ".csv":
        sheet_rows = [(a.replace("№", "N"), b, c) for a, b, c in sheet_rows]
        for g, (a, b, c) in zip(gold_lines, sheet_rows):
            g["source"] = f"{a} {c}".strip()
    rel = f"files/sheet_{seq + 1:03d}{ext}"
    write_sheet_file(files_dir / f"sheet_{seq + 1:03d}{ext}", layout, sheet_rows)
    message = "Добрый день! Заказ в прикрепленном файле."
    return {"id": case_id, "kind": "sheet", "customer_id": customer_id,
            "message": message, "attachment": rel, "content_type": None,
            "gold": gold_lines, "gold_errors": []}


def make_portal_case(rng, ctx, ds, plant, case_id, customer_id, seq, files_dir):
    products = ctx["products"]
    n = rng.randint(5, 15)
    chosen = rng.sample(products, n)
    used = {p.material_id for p in chosen}
    err_count = rng.randint(1, 3)
    err_rows = sorted(rng.sample(range(n), err_count))
    err_codes = rng.sample(list(ERROR_CODES), err_count)
    multi_pool = [p for p in products if p.order_multiple > 1]
    rows: list[list] = []
    for p in chosen:
        qty = p.order_multiple * rng.randint(1, 3)
        rows.append([p.ean13, p.material_id, str(qty), "text"])
    gold_errors = []
    for row_idx, code in zip(err_rows, err_codes):
        p = chosen[row_idx]
        file_row = row_idx + 2
        if code == "not_text":
            rows[row_idx][3] = "number"
        elif code == "bad_ean":
            rows[row_idx][0] = ean_bad(p.ean13)
        elif code == "bad_material_id":
            rows[row_idx][1] = rng.choice(["12345678901234567", "199999999999999999"])
        elif code == "ean_code_mismatch":
            other = rng.choice([q for q in products if q.material_id != p.material_id])
            rows[row_idx][0] = other.ean13
        elif code == "bad_qty":
            rows[row_idx][2] = rng.choice(["10 уп", "0", "0 шт"])
        elif code == "not_multiple":
            if p.order_multiple <= 1 or p.material_id in {r[1] for j, r in enumerate(rows) if j != row_idx}:
                swap = rng.choice([q for q in multi_pool if q.material_id not in used])
                used.discard(p.material_id)
                used.add(swap.material_id)
                chosen[row_idx] = swap
                p = swap
                rows[row_idx][0] = p.ean13
                rows[row_idx][1] = p.material_id
            delta = rng.randint(1, p.order_multiple - 1)
            rows[row_idx][2] = str(p.order_multiple + delta)
        elif code == "column_count":
            rows[row_idx][3] = "extra"
        elif code == "duplicate":
            src = next(j for j in range(n) if j not in err_rows)
            rows[row_idx][0] = rows[src][0]
            rows[row_idx][1] = rows[src][1]
            rows[row_idx][2] = rows[src][2]
            rows[row_idx][3] = rows[src][3]
        gold_errors.append({"row": file_row, "code": code})
    gold_errors.sort(key=lambda e: e["row"])
    rel = f"files/portal_{seq + 1:03d}.xls"
    write_portal_file(files_dir / f"portal_{seq + 1:03d}.xls",
                      [(r[0], r[1], r[2], r[3]) for r in rows])
    message = "Добрый день! Прикрепляю файл для загрузки заказа. Проверьте, пожалуйста."
    return {"id": case_id, "kind": "portal_file", "customer_id": customer_id,
            "message": message, "attachment": rel, "content_type": None,
            "gold": [], "gold_errors": gold_errors}


def generate(out: Path) -> list[dict]:
    out = Path(out)
    files_dir = out / "files"
    files_dir.mkdir(parents=True, exist_ok=True)
    ds = default_source()
    ctx = build_context(ds)
    rng = random.Random(SEED)
    customers = list(ds.customers())
    cases: list[dict] = []
    sheet_seq = portal_seq = photo_seq = 0
    idx = 0
    plan = [("text_ru", N_TEXT_RU), ("text_kk", N_TEXT_KK), ("sheet", N_SHEET),
            ("portal_file", N_PORTAL), ("photo", N_PHOTO), ("long", N_LONG)]
    for kind, total in plan:
        for j in range(total):
            customer = customers[idx % len(customers)]
            idx += 1
            plant = customer.branch_plant
            case_id = f"{kind}_{j + 1:03d}"
            if kind in ("text_ru", "text_kk", "long"):
                kk = kind == "text_kk"
                n = 30 if kind == "long" else rng.randint(3, 12)
                gold_lines = [build_line(rng, ctx, plant, kk)[0] for _ in range(n)]
                message = assemble_text_message(rng, [g["source"] for g in gold_lines], kk)
                cases.append({"id": case_id, "kind": kind, "customer_id": customer.customer_id,
                              "message": message, "attachment": None, "content_type": None,
                              "gold": gold_lines, "gold_errors": []})
            elif kind == "sheet":
                rec = make_sheet_case(rng, ctx, ds, plant, case_id,
                                      customer.customer_id, sheet_seq, files_dir)
                sheet_seq += 1
                cases.append(rec)
            elif kind == "portal_file":
                rec = make_portal_case(rng, ctx, ds, plant, case_id,
                                       customer.customer_id, portal_seq, files_dir)
                portal_seq += 1
                cases.append(rec)
            elif kind == "photo":
                n = rng.randint(4, 8)
                gold_lines = [build_line(rng, ctx, plant, kk=False)[0] for _ in range(n)]
                rel = f"files/photo_{photo_seq + 1:03d}.png"
                render_list([g["source"] for g in gold_lines], files_dir / rel.split("/", 1)[1],
                            rng.randint(0, 10 ** 9))
                photo_seq += 1
                cases.append({"id": case_id, "kind": kind, "customer_id": customer.customer_id,
                              "message": "Здравствуйте! Список от руки на фото.",
                              "attachment": rel, "content_type": "image/png",
                              "gold": gold_lines, "gold_errors": []})
    with open(out / "cases.jsonl", "w", encoding="utf-8") as f:
        for rec in cases:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")
    return cases


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description="Generate deterministic eval cases.")
    parser.add_argument("--out", default="evals/cases")
    args = parser.parse_args(argv)
    generate(Path(args.out))


if __name__ == "__main__":
    main()
